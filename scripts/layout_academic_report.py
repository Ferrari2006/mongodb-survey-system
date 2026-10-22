"""Create a separately saved A4 course report with consistent academic layout.

Run using the bundled Python. Saved user drafts and previous images are not
overwritten. Formatting guidance does not certify submission or correctness.
"""
import base64
import hashlib
import json
import re
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from export_editable_report import Element, NAME, OUTPUT, ROOT, Reader, make_word

SOURCE = OUTPUT / (NAME + "-图表美化编辑版.html")
HTML_OUTPUT = OUTPUT / (NAME + "-排版核对编辑版.html")
WORD_OUTPUT = OUTPUT / (NAME + "-排版核对可编辑.docx")
QA = ROOT / "tmp" / "academic-report"

CAPTIONS = [
    "表1　集合职责与关键字段",
    "表2　索引与对应访问需求",
    "表3　后端模块与主要接口",
    "表4　题目定义与回答的分层校验",
    "表5　课程要求的六类测试及结果",
    "表6　已发现的问题与处理依据",
    "表7　人工反馈与实现调整",
    "表8　AI使用中的问题与修正证据",
]

CSS = r"""
*{box-sizing:border-box}
html{background:#edf0f3}
body{margin:0;color:#111;font:11pt/1.5 'SimSun','Songti SC','Times New Roman',serif}
#report{width:min(100% - 32px,920px);margin:24px auto;padding:40px 54px 48px;background:#fff;outline:none;min-height:90vh}
#report h1,#report h2,#report h3{font-family:'SimHei','Microsoft YaHei',sans-serif;color:#000;break-after:avoid;line-height:1.35;page-break-after:avoid}
#report h1{font-size:19pt;text-align:center;font-weight:700;margin:8pt 0 9pt;border:0;padding:0}
#report h2{font-size:13.5pt;margin:18pt 0 7pt}
#report h3{font-size:11.5pt;margin:12pt 0 6pt}
#report p{margin:0 0 6pt;orphans:3;widows:3}
#report>.repo-link{font-size:9.5pt;margin:0 0 8pt;overflow-wrap:anywhere}
#report>.report-meta{text-align:center;font-size:9.5pt;margin-bottom:16pt;line-height:1.6}
#report ul,#report ol{padding-left:1.9em;margin:4pt 0 8pt}
#report li{padding-left:2pt;margin:0 0 4pt}
#report a{color:#204f70;text-decoration:none}
#report code{font:9.1pt/1.35 Consolas,'SimSun',monospace;overflow-wrap:anywhere}
#report pre{font:8.8pt/1.4 Consolas,'SimSun',monospace;white-space:pre-wrap;overflow-wrap:anywhere;border:1px solid #d9d9d9;background:#f7f8fa;padding:9pt 11pt;margin:8pt 0 9pt;break-inside:avoid}
#report pre code{font-size:inherit;line-height:inherit;overflow-wrap:normal}
.table-block{margin:9pt 0 10pt;break-inside:avoid;page-break-inside:avoid}
#report .table-caption{text-align:center;font-size:10pt;line-height:1.4;margin:0 0 5pt;font-weight:bold;break-after:avoid}
#report table{width:100%;table-layout:fixed;border-collapse:collapse;font-size:9.1pt;line-height:1.35;margin:0}
#report th,#report td{border:1px solid #d9d9d9;padding:5.5pt 6pt;vertical-align:middle;text-align:left;overflow-wrap:anywhere;word-break:normal}
#report th{background:#edf2f6;color:#000;font-weight:bold}
#report th:first-child{width:auto;min-width:0}
#report th code,#report td code{font-size:8.7pt}
#report thead{display:table-header-group}
#report tr{break-inside:avoid;page-break-inside:avoid}
.report-figure{margin:10pt 0 12pt;break-inside:avoid;page-break-inside:avoid}
#report .report-figure p{margin:0}
#report .report-figure img{display:block;width:100%;height:auto;margin:0 auto 6pt;max-height:none}
#report .report-figure p:last-child{font:9pt/1.45 'SimSun','Times New Roman',serif;color:#333;text-align:left}
#editor-bar{position:sticky;top:0;z-index:10;display:flex;gap:10px;flex-wrap:wrap;align-items:center;padding:12px 22px;background:white;border-bottom:1px solid #cbd3dc;font:14px/1.5 'Microsoft YaHei',sans-serif}
#editor-bar button{border:1px solid #718291;background:#f4f7f9;border-radius:4px;padding:7px 13px;color:#17384c;cursor:pointer}
#status{color:#204f70}#edit-help{font-size:12px;color:#4b5b68}
@page{size:A4;margin:18mm 20mm 18mm;@bottom-center{content:counter(page) ' / ' counter(pages);font-family:'Times New Roman',serif;font-size:9pt;color:#555}}
@media(max-width:700px){#report{width:calc(100% - 16px);padding:24px 18px}#report table{font-size:9pt}#report th,#report td{padding:4pt}}
@media print{html,body{background:white}#editor-bar{display:none!important}#report{width:auto;max-width:none;margin:0;padding:0;min-height:0}a{color:#111!important}#report{print-color-adjust:exact;-webkit-print-color-adjust:exact}}
"""


def nodes(node):
    yield node
    for child in node.children:
        if isinstance(child, Element):
            yield from nodes(child)


def data_uri(path, mime):
    return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def configure_word(path):
    doc = Document(path)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.footer_distance = Cm(.9)
    for name, size in [("Normal", 11), ("Title", 19), ("Heading 1", 13.5), ("Heading 2", 11.5)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.line_spacing = 1.35 if name == "Normal" else 1.2
        style.paragraph_format.space_after = Pt(5 if name == "Normal" else 6)
        style.paragraph_format.space_before = Pt(0 if name in ("Normal", "Title") else 12)
    for paragraph in doc.paragraphs:
        pf = paragraph.paragraph_format
        pf.widow_control = True
        text = paragraph.text.strip()
        if text.startswith("表") and re.match(r"表\d", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.keep_with_next = True
            pf.space_before, pf.space_after = Pt(6), Pt(4)
            for r in paragraph.runs:
                r.font.size = Pt(10)
                r.bold = True
        elif text.startswith(("图1", "图2")):
            pf.space_after = Pt(7)
            pf.line_spacing = 1.35
        elif text.startswith("GitHub 仓库"):
            pf.space_after = Pt(8)
            for r in paragraph.runs:r.font.size = Pt(9.5)
        elif text.startswith("课程："):
            pf.space_after = Pt(13)
        if paragraph._p.xpath('.//w:drawing'):
            pf.space_before, pf.space_after = Pt(6), Pt(5)
            pf.line_spacing = 1
        # A JSON example is an intact block rather than loose wrapped lines.
        if text.startswith('{') and '"id"' in text:
            pf.keep_together = True
            pf.line_spacing = 1.25
        if "题目结构节选如下" in text:
            pf.keep_with_next = True
        # The converter's empty post-table paragraphs add an entire blank
        # line. Table padding and caption spacing already separate these blocks.
        if not text and not paragraph._p.xpath('.//w:drawing'):
            paragraph._p.getparent().remove(paragraph._p)
    widths_cm = [[3.1,7.0,6.9], [5.0,12.0], [3.4,13.6], [4.0,13.0],
                 [2.7,6.1,8.2], [4.2,7.1,5.7], [6.0,11.0], [4.8,7.2,5.0]]
    for table, widths in zip(doc.tables, widths_cm):
        for column, width in zip(table.columns, widths):column.width = Cm(width)
        for row_index, row in enumerate(table.rows):
            # These small tables fit within one page; keep rows with the next
            # except the last, while preserving editable native table cells.
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = row_index < len(table.rows)-1
                    p.paragraph_format.line_spacing = 1.2
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.widow_control = True
                    for run in p.runs:run.font.size = Pt(9.1)
    footer = sec.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in ("PAGE", " / ", "NUMPAGES"):
        run = footer.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("555555")
        if text.strip() == "/":run.text = text
        else:
            begin = OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"), "begin")
            instruction = OxmlElement("w:instrText");instruction.text = ' ' + text + ' '
            end = OxmlElement("w:fldChar");end.set(qn("w:fldCharType"), "end")
            run._r.extend([begin, instruction, end])
    doc.save(path)


def main():
    if HTML_OUTPUT.exists() or WORD_OUTPUT.exists():
        raise SystemExit("Layout copy exists; choose a new destination to preserve edits.")
    before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    html = SOURCE.read_text(encoding="utf-8")
    substitutions = {
        "<h2>五、测试、发现的问题与验证结果</h2>": "<h2>五、测试方法与验证结果</h2>",
        "<h2>六、AI协作：需求引导、方案与修正</h2>": "<h2>六、AI协作过程与问题修正</h2>",
        "<h3>3.2 跳转设计：规则保存为数据</h3>": "<h3>3.2 用数据配置跳转规则</h3>",
        "<h3>3.3 校验设计：先检查配置，再检查回答</h3>": "<h3>3.3 题目定义与回答的分层校验</h3>",
        "<code>response.survey_revision</code>": "正式答卷的<code>survey_revision</code>字段",
        "之后的新请求拒绝。": "之后的新提交拒绝；已成功提交的同request_id重试仍返回原回执。",
        'alt="依据2026年10月21日真实测试记录重绘的多选统计与数字题核对"': 'alt="2026年10月21日测试数据重绘。多选有效回答2份，图书馆和食堂各2次即100%，体育设施0次即0%；数字0和8的平均值为4。"',
    }
    for old, new in substitutions.items():
        if html.count(old) != 1:raise ValueError(f"Source changed; review before adjusting: {old}")
        html = html.replace(old, new)
    # The synthetic run counts submitted responses, not independent people.
    image_path = ROOT / "reports" / "images" / "stage1-statistics-academic.svg"
    svg = (ROOT/"reports/images/stage1-statistics-refined.svg").read_text(encoding="utf-8")
    svg = svg.replace("每人可选多项", "每份答卷可选多项")
    if image_path.exists():raise SystemExit("New statistics asset already exists.")
    image_path.write_text(svg, encoding="utf-8")
    i = 0
    def figure(match):
        nonlocal i
        i += 1
        return match[0] if i == 1 else re.sub(r'src="[^"]*"', 'src="'+data_uri(image_path,"image/svg+xml")+'"', match[0])
    html = re.sub(r"<img\b[^>]*>", figure, html)
    assert i == 2
    count = 0
    def table(match):
        nonlocal count
        caption = CAPTIONS[count]
        count += 1
        widths = [[18,41,41],[30,70],[20,80],[24,76],[16,36,48],[25,42,33],[35,65],[28,42,30]][count-1]
        colgroup = "<colgroup>" + "".join(f'<col style="width:{width}%">' for width in widths) + "</colgroup>"
        markup = match[0].replace("<table>", "<table>" + colgroup, 1)
        return '<section class="table-block"><p class="table-caption">' + caption + "</p>" + markup + "</section>"
    html = re.sub(r"<table>[\s\S]*?</table>", table, html)
    assert count == 8
    html = re.sub(r"<style>[\s\S]*?</style>", "<style>" + CSS + "</style>", html, count=1)
    html = html.replace("<p>GitHub 仓库：", '<p class="repo-link">GitHub 仓库：', 1)
    html = html.replace("<p>课程：", '<p class="report-meta">课程：', 1)
    html = html.replace("第一阶段报告 图表美化编辑版", "第一阶段报告 排版核对编辑版")
    html = html.replace("<strong>报告图表美化版</strong>", "<strong>报告排版核对版</strong>")
    html = html.replace(NAME + "-图表美化-我的修改.html", NAME + "-排版核对-我的修改.html")
    HTML_OUTPUT.write_text(html, encoding="utf-8")
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == before
    QA.mkdir(parents=True, exist_ok=True)
    (QA/"source.json").write_text(json.dumps({"source":SOURCE.name,"source_sha256":before,"source_unchanged":True,
                                             "html":HTML_OUTPUT.name,"table_captions":CAPTIONS},ensure_ascii=False,indent=2),encoding="utf-8")
    print("New HTML saved. Rasterize the new SVG, then call build_word().")


def build_word():
    if WORD_OUTPUT.exists():raise SystemExit("Word copy exists; preserve it.")
    reader=Reader();reader.feed(HTML_OUTPUT.read_text(encoding="utf-8"))
    report=next(n for n in nodes(reader.root) if n.attrs.get("id")=="report")
    for node in nodes(report):
        if node.tag=="img":
            stem="stage1-data-model-refined" if "五集合" in node.attrs.get("alt","") else "stage1-statistics-academic"
            node.attrs["src"]=data_uri(ROOT/"reports/images"/(stem+".png"),"image/png")
            node.attrs["data-word-width-cm"]="16.8"
    # The native Word converter does not need HTML-only column metadata.
    for node in list(nodes(report)):
        node.children=[c for c in node.children if not (isinstance(c,Element) and c.tag=="colgroup")]
    make_word(report,WORD_OUTPUT)
    configure_word(WORD_OUTPUT)
    print(WORD_OUTPUT.name.encode('ascii','backslashreplace').decode())


if __name__=="__main__":
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--word',action='store_true');args=p.parse_args()
    build_word() if args.word else main()
