"""Create new stage-two submission artifacts without modifying saved drafts."""
import hashlib
import json
from pathlib import Path
import sys
import build_stage2_draft as template

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'第二阶段/第二阶段完成报告.md'
OUT=ROOT/'第二阶段/提交材料'
NAME='2026_ECNU_PJ1_2_温正泽'

def main():
    OUT.mkdir(exist_ok=True)
    if (OUT/(NAME+'-可编辑.docx')).exists() and '--refresh' not in sys.argv:
        raise SystemExit('Existing submission preserved. Pass --refresh only for generated artifacts, not saved user edits.')
    observed=json.loads((ROOT/'tests/results/stage2-latest.json').read_text(encoding='utf-8'))
    assert observed['successful'] and observed['tests_run']==52
    for name in ('stage2-acceptance.json','stage2-reproduction.json'):
        assert json.loads((ROOT/'tests/results'/name).read_text(encoding='utf-8'))['successful']
    template.SOURCE=SOURCE;template.OUT=OUT;template.NAME=NAME
    blocks=template.blocks();template.build_html(blocks);template.build_word(blocks)
    from docx import Document
    from docx.shared import Pt,Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document(OUT/(NAME+'-可编辑.docx'))
    def set_fonts(paragraph, heading=False):
        for run in paragraph.runs:
            fonts=run._element.get_or_add_rPr().get_or_add_rFonts()
            for attr in list(fonts.attrib):
                if 'theme' in attr.lower():del fonts.attrib[attr]
            for attr in ('ascii','hAnsi','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
            fonts.set(qn('w:eastAsia'),'黑体' if heading else '宋体')
            if heading:run.bold=False
    for paragraph in doc.paragraphs:
        style=paragraph.style.name
        set_fonts(paragraph,style in ('Title','Heading 1','Heading 2'))
        if style=='Title':
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:run.font.size=Pt(16)
        elif style in ('Heading 1','Heading 2'):
            paragraph.paragraph_format.space_before=Pt(12 if style=='Heading 1' else 8)
            paragraph.paragraph_format.space_after=Pt(5)
            for run in paragraph.runs:run.font.size=Pt(13.5 if style=='Heading 1' else 12)
        elif paragraph.text and not paragraph.text.startswith(('GitHub','温正泽')):
            caption=paragraph.text.startswith(('图1 ','图2 ','图3 ','图4 ','表1 ','表2 ','表3 ','表4 '))
            for run in paragraph.runs:run.font.size=Pt(10 if caption else 11.5)
            if caption:paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.paragraph_format.line_spacing=1.45
                paragraph.paragraph_format.first_line_indent=Pt(23)
                paragraph.paragraph_format.space_after=Pt(2)
        if paragraph.text.startswith('附录 页面'):paragraph.paragraph_format.page_break_before=False
    widths=[(3.2,13.8),(4.2,6.3,6.5),(3.4,6.2,7.4),(5,8,4)]
    for table_index,(table,columns) in enumerate(zip(doc.tables,widths)):
        for column,width in zip(table.columns,columns):column.width=Cm(width)
        for i,row in enumerate(table.rows):
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for cell,width in zip(row.cells,columns):
                cell.width=Cm(width)
                for paragraph in cell.paragraphs:
                    set_fonts(paragraph,heading=i==0)
                    paragraph.paragraph_format.line_spacing=1.25
                    paragraph.paragraph_format.space_before=Pt(3)
                    paragraph.paragraph_format.space_after=Pt(3)
                    for run in paragraph.runs:run.font.size=Pt(10.2)
                    if i==0 or (table_index==1 and i<len(table.rows)-1):
                        paragraph.paragraph_format.keep_with_next=True
    for section in doc.sections:
        for container in (section.header,section.footer):
            for paragraph in container.paragraphs:set_fonts(paragraph)
    doc.save(OUT/(NAME+'-可编辑.docx'))
    p=OUT/(NAME+'-可编辑.html')
    html=p.read_text(encoding='utf-8').replace('报告草稿编辑版','报告编辑版').replace('第二阶段报告草稿','第二阶段项目报告')
    html=html.replace('</style>','''#report{font-family:"Times New Roman","SimSun",serif;}
#report h1,#report h2,#report h3,#report th{font-family:"Times New Roman","SimHei",sans-serif;font-weight:400;}
#report p{font-size:11.5pt;line-height:1.45;text-indent:2em;}
#report h1{font-size:16pt;text-align:center;}#report h2{font-size:13.5pt;}#report h3{font-size:12pt;}
#report table{font-size:10.2pt;line-height:1.25;}#report table p{font-size:inherit;text-indent:0;}
#report .table-caption{font-size:10pt;text-align:center;text-indent:0;}
#report>p:first-child,#report .report-meta{text-indent:0;font-size:9.5pt;}
</style>''',1)
    import re
    column_sets=iter(widths)
    def add_columns(match):
        columns=next(column_sets)
        return '<table><colgroup>'+''.join(f'<col style="width:{width/17*100:.3f}%">' for width in columns)+'</colgroup>'
    html=re.sub(r'<table>',add_columns,html)
    p.write_text(html,encoding='utf-8')
    manifest={'stage':'stage2-submission','source':SOURCE.relative_to(ROOT).as_posix(),
        'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'recorded_tests':observed['tests_run'],
        'files':[NAME+'-可编辑.html',NAME+'-可编辑.docx']}
    (OUT/'submission-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('Stage-two submission HTML and Word created; previous drafts retained')

if __name__=='__main__':main()
