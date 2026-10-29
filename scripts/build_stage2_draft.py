"""Build a stage-two review draft from the retained stage-one visual template.

No submitted stage-one file is written. Rerunning requires explicit --refresh.
"""
import base64
from copy import deepcopy
import hashlib
import html
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'第二阶段/报告草稿'
SOURCE=ROOT/'第二阶段/第二阶段完成报告草稿.md'
REF=ROOT/'第一阶段/提交材料/2026_ECNU_PJ1_1_温正泽-AI过程修订可编辑.docx'
HTML_REF=ROOT/'第一阶段/提交材料/2026_ECNU_PJ1_1_温正泽-AI过程修订编辑版.html'
NAME='2026_ECNU_PJ1_2_温正泽-草稿'

def svg(title,parts,height=400):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="960" height="{height}" viewBox="0 0 960 {height}" role="img"><title>{title}</title><defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0 0L8 4L0 8" fill="#456878"/></marker></defs><rect width="960" height="{height}" fill="white"/><g font-family="Microsoft YaHei,SimSun,sans-serif" fill="#243b47">'+''.join(parts)+'</g></svg>'

def box(x,y,w,h,title,lines,color):
    return f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="12" fill="{color}" stroke="#9cb5bd"/><text x="{x+w/2}" y="{y+34}" text-anchor="middle" font-size="24" font-weight="bold">{title}</text>'+''.join(f'<text x="{x+w/2}" y="{y+66+i*28}" text-anchor="middle" font-size="21">{line}</text>' for i,line in enumerate(lines))

def arrow(x1,y1,x2,y2,dashed=False):
    return f'<path d="M{x1} {y1} L{x2} {y2}" stroke="#456878" stroke-width="2.5" fill="none" marker-end="url(#arrow)"'+(' stroke-dasharray="7 5"' if dashed else '')+'/>'

def figures():
    target=ROOT/'第二阶段/images';target.mkdir(exist_ok=True)
    p=[box(20,25,270,160,'题库与题目',['question_banks','library_questions','工作稿与历史版本'],'#e5f1ed'),
       box(345,25,270,160,'问卷内副本',['surveys.questions','来源题与版本ID','独立题号与本卷跳转'],'#e8eef6'),
       box(670,25,270,160,'正式答卷',['responses','保留原问卷题号','仅统计已提交内容'],'#f9eee6'),
       arrow(290,100,345,100),arrow(615,100,670,100),
       '<text x="317" y="215" text-anchor="middle" font-size="20">指定版本复制</text>',
       box(160,265,640,108,'按本人问卷汇总',['来源版本 + 实际定义摘要；不同含义不混算'],'#eeeaf5'),
       arrow(480,185,480,265),arrow(805,185,700,265),
       '<text x="480" y="410" text-anchor="middle" font-size="21">源题后续修改不会自动覆盖已导入的问卷</text>']
    (target/'stage2-data-flow.svg').write_text(svg('题库来源与问卷快照关系',p,440),encoding='utf-8')
    p=[box(25,20,250,110,'第1版',['初始内容'],'#e5f1ed'),box(355,20,250,110,'第2版',['父版本为第1版'],'#e8eef6'),box(685,20,250,110,'第3版',['父版本为第2版'],'#e8eef6'),
       box(355,235,250,110,'第4版',['从第1版分支'],'#eeeaf5'),box(685,235,250,110,'第5版',['父版本为第4版','恢复来源为第1版'],'#f9eee6'),
       arrow(275,75,355,75),arrow(605,75,685,75),arrow(150,130,355,290),arrow(605,290,685,290),
       '<text x="480" y="175" text-anchor="middle" font-size="21">每版保留内容与父版本；恢复追加记录</text>',
       '<text x="480" y="395" text-anchor="middle" font-size="20">版本关系示意；各版本可分别导入不同问卷</text>']
    (target/'stage2-version-flow.svg').write_text(svg('题目版本分支与恢复示意',p,425),encoding='utf-8')

def blocks():
    lines=SOURCE.read_text(encoding='utf-8').splitlines();i=0;result=[]
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|'):
            rows=[]
            while True:
                if not re.fullmatch(r'[| :\-]+',line):rows.append([s.strip() for s in line.strip('|').split('|')])
                if i>=len(lines) or not lines[i].startswith('|'):break
                line=lines[i].strip();i+=1
            result.append(('table',rows));continue
        match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if match:result.append(('image',(match[1],match[2])));continue
        match=re.match(r'^(#{1,3}) (.+)',line)
        if match:result.append(('h'+str(len(match[1])),match[2]));continue
        result.append(('p',line))
    return result

def inline(value):return re.sub(r'`([^`]+)`',r'<code>\1</code>',html.escape(value))

def build_html(data):
    source=HTML_REF.read_text(encoding='utf-8').replace('第一阶段','第二阶段').replace('Stage 1','Stage 2').replace('2026_ECNU_PJ1_1_温正泽',NAME)
    source=source.replace('AI过程修订编辑版','报告草稿编辑版').replace('报告AI过程修订版','第二阶段报告草稿')
    body=['<div id="page-header-preview" contenteditable="false"><span>数据管理系统课程设计 · 第二阶段</span><span>温正泽 10244304404</span></div>']
    for index,(kind,value) in enumerate(data):
        if kind=='table':
            rows=['<tr>'+''.join(f'<{"th" if j==0 else "td"}>{inline(cell)}</{"th" if j==0 else "td"}>' for cell in row)+'</tr>' for j,row in enumerate(value)]
            body.append('<section class="table-block"><table><thead>'+rows[0]+'</thead><tbody>'+''.join(rows[1:])+'</tbody></table></section>')
        elif kind=='image':
            alt,path=value;encoded=base64.b64encode((ROOT/'第二阶段'/path).read_bytes()).decode()
            mime='image/png' if Path(path).suffix.lower()=='.png' else 'image/svg+xml'
            body.append(f'<figure class="report-figure"><img src="data:{mime};base64,{encoded}" alt="{html.escape(alt)}" style="width:100%" contenteditable="false"></figure>')
        elif kind=='p' and index==0:
            url=value.split('：',1)[1];body.append(f'<p class="repo-link">GitHub 仓库：<a href="{url}">{url}</a></p>')
        else:
            cls=' class="report-meta"' if kind=='p' and value.startswith('温正泽') else ' class="table-caption"' if kind=='p' and re.match(r'[表图]\d',value) else ''
            body.append(f'<{kind}{cls}>{inline(value)}</{kind}>')
    result=re.sub(r'(<main\b[^>]*>)[\s\S]*?(</main>)',lambda m:m.group(1)+''.join(body)+m.group(2),source,count=1)
    (OUT/(NAME+'-可编辑.html')).write_text(result,encoding='utf-8')

def build_word(data):
    from docx import Document
    from docx.shared import Cm,Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document(REF)
    normal=next(p for p in doc.paragraphs if p.paragraph_format.first_line_indent)
    normal_pr=deepcopy(normal._p.pPr)
    table_template=deepcopy(doc.tables[0]._tbl.tblPr)
    for child in list(doc._element.body):
        if child.tag!=qn('w:sectPr'):doc._element.body.remove(child)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:run.text=run.text.replace('第一阶段','第二阶段')
    for index,(kind,value) in enumerate(data):
        if kind=='table':
            table=doc.add_table(rows=len(value),cols=len(value[0]));table._tbl.replace(table._tbl.tblPr,deepcopy(table_template))
            table.autofit=False
            for j,row in enumerate(value):
                for k,text in enumerate(row):
                    cell=table.cell(j,k);cell.text=text
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.35
                        for run in p.runs:run.font.size=Pt(9.3);run.bold=j==0
                    if j==0:
                        borders=OxmlElement('w:tcBorders');edge=OxmlElement('w:bottom');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');borders.append(edge);cell._tc.get_or_add_tcPr().append(borders)
                if j==0:
                    repeat=OxmlElement('w:tblHeader');table.rows[j]._tr.get_or_add_trPr().append(repeat)
            continue
        if kind=='image':
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_together=True
            p.add_run().add_picture(str((ROOT/'第二阶段'/value[1]).with_suffix('.png')),width=Cm(17));continue
        style={'h1':'Title','h2':'Heading 1','h3':'Heading 2'}.get(kind,'Normal')
        p=doc.add_paragraph(style=style)
        if kind=='p' and index>2 and not re.match(r'[表图]\d',value):
            p._p.replace(p._p.pPr,deepcopy(normal_pr)) if p._p.pPr is not None else p._p.insert(0,deepcopy(normal_pr))
        p.add_run(value.replace('`',''))
        if kind=='p' and index>2 and not re.match(r'[表图]\d',value):p.paragraph_format.space_after=Pt(3)
        if kind in ('h2','h3'):
            p.paragraph_format.space_before=Pt(10 if kind=='h2' else 9);p.paragraph_format.space_after=Pt(5)
        if kind=='h1':
            p.paragraph_format.first_line_indent=Pt(0)
        if kind=='p' and (index==0 or value.startswith('温正泽')):
            p.paragraph_format.first_line_indent=Pt(0)
            for r in p.runs:r.font.size=Pt(9.5)
        if kind=='p' and re.match(r'[表图]\d',value):
            p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.size=Pt(9.5)
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    doc.save(OUT/(NAME+'-可编辑.docx'))
    # Body, image relationships, header text and field-refresh setting are intended slots.
    preserved=['word/styles.xml','word/numbering.xml','word/theme/theme1.xml']
    with ZipFile(REF) as a,ZipFile(OUT/(NAME+'-可编辑.docx')) as b:
        for name in preserved:assert a.read(name)==b.read(name),f'Unexpected template change: {name}'

def main():
    if len(sys.argv)>1 and sys.argv[1]=='--figures':figures();return
    OUT.mkdir(exist_ok=True)
    if (OUT/(NAME+'-可编辑.html')).exists() and '--refresh' not in sys.argv:raise SystemExit('Existing draft preserved. Use --refresh only for generated drafts, never user-edited copies.')
    observed=json.loads((ROOT/'tests/results/stage2-latest.json').read_text(encoding='utf-8'))
    assert observed['successful'] and observed['tests_run']==52
    data=blocks();build_html(data);build_word(data)
    manifest={'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),
        'recorded_tests':observed['tests_run'],'stage':'stage2-review-draft','files':[NAME+'-可编辑.html',NAME+'-可编辑.docx']}
    (OUT/'draft-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('Stage-two HTML and Word review drafts created; retained reference unchanged.')

if __name__=='__main__':main()
