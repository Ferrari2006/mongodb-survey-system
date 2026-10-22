"""Replace the data diagram and add a source-backed layered architecture figure."""
import base64
from pathlib import Path
from export_editable_report import ROOT,OUTPUT,NAME,Element,make_word
from layout_academic_report import nodes,configure_word
from finalize_stage1_report import serialize,fragment
import align_stage1_delivery as align

EDITOR=OUTPUT/(NAME+'-分层图示编辑版.html')
WORD=OUTPUT/(NAME+'-分层图示可编辑.docx')


def uri(name):
    return 'data:image/svg+xml;base64,'+base64.b64encode((ROOT/'reports/images'/name).read_bytes()).decode()


def build():
    import re
    if EDITOR.exists():raise SystemExit('Preserve existing diagram editing copy.')
    html=(OUTPUT/(NAME+'-字体优化编辑版.html')).read_text(encoding='utf-8')
    report=align.parse_report(html)
    images=[n for n in nodes(report) if n.tag=='img']
    images[0].attrs.update({'src':uri('stage1-data-model-layered.svg'),'data-source-file':'stage1-data-model-layered.svg','alt':'五集合分层关系与题目嵌入结构'})
    images[1].attrs['data-source-file']='stage1-statistics-academic.svg'
    for n in nodes(report):
        n.children=[c.replace('图2','图3') if isinstance(c,str) else c for c in n.children]
    heading=next(n for n in report.children if isinstance(n,Element) and n.tag=='h3' and n.text().startswith('3.1'))
    figure=Element('figure',{'class':'report-figure architecture-figure'})
    p=Element('p');image=Element('img',{'src':uri('stage1-architecture-layered.svg'),'alt':'系统分层架构：原生网页、Flask入口、后端规则模块、PyMongo与MongoDB五集合','data-source-file':'stage1-architecture-layered.svg','contenteditable':'false'});p.children=[image]
    caption=Element('p');caption.children=['图2　系统分层架构。按当前源码中的职责组织，规则与校验在后端执行；数据访问为PyMongo调用方式，并非额外独立服务。']
    figure.children=[p,caption];report.children.insert(report.children.index(heading)+1,figure)
    report.children.insert(0,Element('div',{'id':'page-header-preview','contenteditable':'false'}));report.children[0].children=fragment('<span>数据管理系统课程设计 · 第一阶段</span><span>温正泽 10244304404</span>')
    html=re.sub(r'(<main\b[^>]*>)[\s\S]*?(</main>)',lambda m:m.group(1)+''.join(serialize(c) for c in report.children)+m.group(2),html,count=1)
    html=re.sub(r'<section class="table-block"><p class="table-caption">(表[367])',r'<section class="table-block keep-table"><p class="table-caption">\1',html)
    html=html.replace('</style>','#report .architecture-figure img{width:100%}#report .architecture-figure{break-inside:avoid;page-break-inside:avoid}\n</style>',1)
    html=html.replace('字体优化版','分层图示版').replace('字体优化编辑版','分层图示编辑版').replace(NAME+'-字体优化-我的修改.html',NAME+'-分层图示-我的修改.html')
    EDITOR.write_text(html,encoding='utf-8');print('New report copy with 3 figures; original saved.')


def word():
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if WORD.exists():raise SystemExit('Preserve existing Word.')
    report=align.parse_report(EDITOR.read_text(encoding='utf-8'))
    for n in nodes(report):
        if n.tag=='img':
            png=(ROOT/'reports/images'/n.attrs['data-source-file']).with_suffix('.png')
            n.attrs['src']='data:image/png;base64,'+base64.b64encode(png.read_bytes()).decode();n.attrs['data-word-width-cm']='16.8'
    make_word(report,WORD);configure_word(WORD)
    doc=Document(WORD);sec=doc.sections[0];sec.top_margin=Cm(2.1);sec.header_distance=Cm(.75)
    for p in doc.paragraphs:
        if p.style.name=='Normal' and p.text and not p.text.startswith(('课程：','GitHub 仓库','图1','图2','图3','表','报告辅助工具','{')) and not p._p.xpath('.//w:drawing'):
            p.paragraph_format.first_line_indent=Pt(21);p.paragraph_format.line_spacing=1.5
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].font.name='Times New Roman'
    p=sec.header.paragraphs[0];p.clear();p.style='Normal';p.paragraph_format.tab_stops.clear_all();p.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT)
    r=p.add_run('数据管理系统课程设计 · 第一阶段\t温正泽 10244304404');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')
    b=OxmlElement('w:pBdr');edge=OxmlElement('w:bottom')
    for k,v in [('val','single'),('sz','4'),('color','555555'),('space','4')]:edge.set(qn('w:'+k),v)
    b.append(edge);p._p.get_or_add_pPr().append(b)
    for t in doc.tables:
        pr=t._tbl.tblPr
        for b in list(pr.findall(qn('w:tblBorders'))):pr.remove(b)
        b=OxmlElement('w:tblBorders')
        for k in ['top','bottom','left','right','insideH','insideV']:
            e=OxmlElement('w:'+k);e.set(qn('w:val'),'single' if k in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'6');b.append(e)
        pr.append(b)
        for i,row in enumerate(t.rows):
            for cell in row.cells:
                tc=cell._tc.get_or_add_tcPr()
                for sh in list(tc.findall(qn('w:shd'))):tc.remove(sh)
                if i==0:
                    b=OxmlElement('w:tcBorders');e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e);tc.append(b)
    doc.save(WORD)


if __name__=='__main__':
    import sys
    if '--word' in sys.argv:word()
    elif '--sync' in sys.argv:align.WORD=WORD;align.sync(EDITOR)
    else:build()
