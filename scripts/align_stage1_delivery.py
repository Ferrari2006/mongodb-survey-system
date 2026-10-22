"""Keep a saved editing source, canonical reports and repository evidence aligned."""
import argparse
import ast
import hashlib
import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from export_editable_report import ROOT, NAME, OUTPUT, Reader, Element, make_word
from layout_academic_report import nodes, data_uri, configure_word

EDITOR = OUTPUT/(NAME+'-仓库一致编辑版.html')
WORD = OUTPUT/(NAME+'-仓库一致可编辑.docx')
QA = ROOT/'tmp/repository-alignment'
LEFT = '数据管理系统课程设计 · Project 1 / Stage 1'
RIGHT = 'MongoDB Survey System'
HEADER_CSS = '''
#page-header-preview{display:flex;justify-content:space-between;gap:15px;border-bottom:1px solid #888;padding:0 0 5px;margin:0 0 20px;color:#555;font:9pt 'Times New Roman','SimSun',serif}
@page{margin-top:22mm;@top-left{content:"数据管理系统课程设计 · Project 1 / Stage 1";font:8pt 'Times New Roman','SimSun',serif;color:#555;width:70%;text-align:left;vertical-align:bottom;padding-bottom:2mm;border-bottom:.5pt solid #888;margin-bottom:4mm}@top-right{content:"MongoDB Survey System";font:8pt 'Times New Roman',serif;color:#555;width:30%;text-align:right;vertical-align:bottom;padding-bottom:2mm;border-bottom:.5pt solid #888;margin-bottom:4mm}}
@media print{#page-header-preview{display:none!important}}
@media print{#report p{break-inside:avoid}}
'''


def parse_report(html):
    reader=Reader();reader.feed(html)
    report=next(n for n in nodes(reader.root) if n.attrs.get('id')=='report')
    for n in list(nodes(report)):
        n.children=[c for c in n.children if not(isinstance(c,Element) and (c.tag=='colgroup' or c.attrs.get('id')=='page-header-preview'))]
    return report


def body_hash(html):
    return hashlib.sha256(re.sub(r'\s+','',parse_report(html).text()).encode()).hexdigest()


def inline(n):
    if isinstance(n,str):return n
    value=''.join(inline(c) for c in n.children)
    if n.tag in ('b','strong'):return '**'+value+'**'
    if n.tag in ('i','em'):return '*'+value+'*'
    if n.tag=='code':return '`'+n.text()+'`'
    if n.tag=='a':return '['+value+']('+n.attrs['href']+')'
    if n.tag=='br':return '<br>'
    return value


def markdown(report):
    def block(n):
        if isinstance(n,str):return ''
        if n.tag in ('h1','h2','h3'):return '#'*int(n.tag[1])+' '+inline(n)+'\n\n'
        if n.tag=='p':
            pics=[c for c in n.children if isinstance(c,Element) and c.tag=='img']
            if pics:
                pic=pics[0];stem=pic.attrs.get('data-source-file') or ('stage1-data-model-refined.svg' if '五集合' in pic.attrs.get('alt','') else 'stage1-statistics-academic.svg')
                return '!['+pic.attrs.get('alt','')+'](images/'+stem+')\n\n'
            return inline(n)+'\n\n'
        if n.tag=='pre':return '```json\n'+n.text().rstrip()+'\n```\n\n'
        if n.tag in ('ul','ol'):
            return ''.join('- '+inline(c)+'\n' for c in n.children if isinstance(c,Element) and c.tag=='li')+'\n'
        if n.tag=='table':
            rows=[r for r in nodes(n) if r.tag=='tr'];text=[]
            for i,row in enumerate(rows):
                cells=[c for c in row.children if isinstance(c,Element) and c.tag in ('th','td')]
                text.append('| '+' | '.join(inline(c).replace('|','\\|').replace('\n',' ') for c in cells)+' |')
                if i==0:text.append('|'+'---|'*len(cells))
            return '\n'.join(text)+'\n\n'
        return ''.join(block(c) for c in n.children)
    return ''.join(block(c) for c in report.children).rstrip()+'\n'


def create(source):
    if EDITOR.exists():raise SystemExit('Editor exists; preserve edits.')
    html=source.read_text(encoding='utf-8')
    html=html.replace('</style>',HEADER_CSS+'</style>',1)
    html=re.sub(r'(<main\b[^>]*>)',r'\1<div id="page-header-preview" contenteditable="false"><span>'+LEFT+'</span><span>'+RIGHT+'</span></div>',html,count=1)
    html=html.replace('报告科学复核版','报告仓库一致版').replace('第一阶段报告 科学复核优化版','第一阶段报告 仓库一致编辑版')
    html=html.replace(NAME+'-科学复核-我的修改.html',NAME+'-仓库一致-我的修改.html')
    EDITOR.write_text(html,encoding='utf-8')
    QA.mkdir(parents=True,exist_ok=True)
    (QA/'input.json').write_text(json.dumps({'source':str(source.relative_to(ROOT)),'sha256':hashlib.sha256(source.read_bytes()).hexdigest()},ensure_ascii=False,indent=2),encoding='utf-8')
    assert body_hash(html)==body_hash(source.read_text(encoding='utf-8'))
    print('Created separate editor with running header.')


def word():
    if WORD.exists():raise SystemExit('Word exists; preserve edits.')
    report=parse_report(EDITOR.read_text(encoding='utf-8'))
    for n in nodes(report):
        if n.tag=='img':
            stem='stage1-data-model-refined' if '五集合' in n.attrs.get('alt','') else 'stage1-statistics-academic'
            n.attrs['src']=data_uri(ROOT/'reports/images'/(stem+'.png'),'image/png')
            n.attrs['data-word-width-cm']='16.8'
    make_word(report,WORD);configure_word(WORD)
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document(WORD);s=d.sections[0];s.top_margin=Cm(2.0);s.header_distance=Cm(.75)
    p=s.header.paragraphs[0];p.clear();p.style='Normal';p.paragraph_format.tab_stops.clear_all();p.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after=Pt(5)
    run=p.add_run(LEFT+'\t'+RIGHT);run.font.name='Times New Roman';run.font.size=Pt(8)
    run.font.color.rgb=RGBColor.from_string('555555');run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    borders=OxmlElement('w:pBdr');edge=OxmlElement('w:bottom')
    for key,val in [('val','single'),('sz','4'),('space','4'),('color','888888')]:edge.set(qn('w:'+key),val)
    borders.append(edge);p._p.get_or_add_pPr().append(borders)
    for p in d.paragraphs:
        if p.text.startswith('报告辅助工具：'):
            p.paragraph_format.line_spacing=1.25
            for run in p.runs:run.font.size=Pt(9)
    d.save(WORD);print('Created separate Word with running header.')


def sync(source):
    html=source.read_text(encoding='utf-8');report=parse_report(html)
    # Keep toolbar and scripts in the editing copy; canonical report is print-only.
    content=re.search(r'<main\b[^>]*>([\s\S]*?)</main>',html).group(1)
    styles=re.search(r'<style>([\s\S]*?)</style>',html).group(1)
    canonical='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>第一阶段完成报告</title><style>'+styles+'</style></head><body><main id="report">'+content+'</main></body></html>'
    destinations=[ROOT/'reports/第一阶段完成报告.md',ROOT/'reports/第一阶段完成报告.html',OUTPUT/(NAME+'.html')]
    archive=ROOT/'第一阶段/历史材料'/('报告统一入口前_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
    archive.mkdir(parents=True,exist_ok=False)
    for file in destinations:
        if file.exists():
            dest=archive/('reports_' + file.name if file.parent.name=='reports' else file.name)
            shutil.copy2(file,dest)
            assert hashlib.sha256(dest.read_bytes()).hexdigest()==hashlib.sha256(file.read_bytes()).hexdigest()
    destinations[0].write_text(markdown(report),encoding='utf-8')
    for file in destinations[1:]:file.write_text(canonical,encoding='utf-8')
    evidence=['backend/app.py','backend/database.py','backend/logic.py','backend/validation.py','backend/statistics.py',
              'database/schema.json','tests/results/latest.json','docs/API说明.md','docs/数据库与系统设计.md']
    hashes={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in evidence}
    manifest={'schema_version':1,'editor':str(source.relative_to(ROOT)).replace('\\','/'),'word':str(WORD.relative_to(ROOT)).replace('\\','/'),
              'markdown':'reports/第一阶段完成报告.md','printable':'reports/第一阶段完成报告.html',
              'submission':str(destinations[2].relative_to(ROOT)).replace('\\','/'),'normalized_body_sha256':body_hash(canonical),
              'files_sha256':{str(f.relative_to(ROOT)).replace('\\','/'):hashlib.sha256(f.read_bytes()).hexdigest() for f in destinations+[source,WORD]},
              'evidence_sha256':hashes,'recorded_tests':38,'backend_tests_rerun':False}
    (ROOT/'reports/stage1-report-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Canonical Markdown, printable report, submission copy and evidence manifest aligned.')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--word',action='store_true');parser.add_argument('--sync-only',action='store_true');parser.add_argument('--source',type=Path)
    args=parser.parse_args()
    if args.word:word()
    elif args.sync_only:sync(args.source or EDITOR)
    else:create(args.source or OUTPUT/(NAME+'-科学复核优化版.html'))
