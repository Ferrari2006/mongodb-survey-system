"""Create native editable Word content and a standalone browser editing copy.

Run with the Codex bundled Python. Neither output overwrites the Markdown or
the printable HTML. Pass --force only after intentionally backing up user edits.
"""
import argparse
import base64
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NAME = "2026_ECNU_PJ1_1_温正泽"
OUTPUT = ROOT / "第一阶段" / "提交材料"


class Element:
    def __init__(self, tag, attrs=None):
        self.tag, self.attrs, self.children = tag, dict(attrs or []), []

    def text(self):
        return "".join(c if isinstance(c, str) else c.text() for c in self.children)


class Reader(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Element("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Element(tag, attrs)
        self.stack[-1].children.append(node)
        if tag not in ("img", "br", "hr", "meta", "link", "input"):
            self.stack.append(node)

    def handle_endtag(self, tag):
        for index in range(len(self.stack)-1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                return

    def handle_data(self, text):
        self.stack[-1].children.append(text)


def font(run, size=None, east="宋体", latin="Calibri"):
    run.font.name = latin
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size:
        run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)


def inlines(paragraph, element, bold=False, mono=False):
    for child in element.children:
        if isinstance(child, str):
            run = paragraph.add_run(child)
            run.bold = bold
            font(run, 9 if mono else None, latin="Consolas" if mono else "Calibri")
        elif child.tag == "a":
            link = OxmlElement("w:hyperlink")
            rel = paragraph.part.relate_to(child.attrs["href"], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
            link.set(qn("r:id"), rel)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "245287"); props.append(color)
            run.append(props)
            text = OxmlElement("w:t"); text.text = child.text(); run.append(text)
            link.append(run); paragraph._p.append(link)
        elif child.tag == "br":
            paragraph.add_run().add_break()
        elif child.tag != "img":
            inlines(paragraph, child, bold or child.tag in ("strong", "b"), mono or child.tag == "code")


def make_word(body, destination):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for name, size in (("Normal", 10.5), ("Title", 19), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体" if name == "Normal" else "黑体")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.35
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    doc.core_properties.title = "基于 MongoDB 的问卷系统第一阶段报告"
    doc.core_properties.author = "温正泽"
    doc.core_properties.subject = "数据管理系统课程第一阶段报告"

    def image(element):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        encoded = element.attrs["src"].split(",", 1)[1]
        if "image/svg+xml" in element.attrs["src"]:
            data = (ROOT / "reports" / "images" / "stage1-data-model-word.png").read_bytes()
        else:
            data = base64.b64decode(encoded)
        picture = paragraph.add_run().add_picture(io.BytesIO(data), width=Cm(16.9 if "image/svg+xml" in element.attrs["src"] else 14.8))
        picture._inline.docPr.set("descr", element.attrs.get("alt", ""))
        paragraph.paragraph_format.keep_with_next = True

    def block(element):
        if isinstance(element, str):
            return
        tag = element.tag
        if tag in ("h1", "h2", "h3"):
            p = doc.add_paragraph(style={"h1":"Title", "h2":"Heading 1", "h3":"Heading 2"}[tag])
            inlines(p, element)
            for run in p.runs:
                font(run, east="黑体")
                run.bold = True
            if tag == "h1": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tag == "p":
            pictures = [c for c in element.children if isinstance(c, Element) and c.tag == "img"]
            if pictures:
                for pic in pictures: image(pic)
                return
            p = doc.add_paragraph()
            inlines(p, element)
            if element.text().startswith("课程："):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs: font(r, 9)
            if element.text().startswith(("图1", "图2")):
                for r in p.runs: font(r, 9)
        elif tag in ("ul", "ol"):
            for child in element.children:
                if isinstance(child, Element) and child.tag == "li":
                    p = doc.add_paragraph(style="List Bullet" if tag == "ul" else "List Number")
                    inlines(p, child)
        elif tag == "pre":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            for line in element.text().rstrip().splitlines():
                r = p.add_run(line + "\n"); font(r, 8.5, latin="Consolas")
        elif tag == "table":
            rows = []
            def collect(node):
                if node.tag == "tr": rows.append(node)
                else:
                    for child in node.children:
                        if isinstance(child, Element): collect(child)
            collect(element)
            cols = len([c for c in rows[0].children if isinstance(c, Element)])
            table = doc.add_table(rows=0, cols=cols)
            table.autofit = False
            widths = [3.4, 7.1, 6.9] if cols == 3 else [4.6, 12.8]
            for col, width in zip(table.columns, widths): col.width = Cm(width)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                item = OxmlElement("w:" + edge)
                for key, value in (("val","single"),("sz","4"),("color","D9D9D9")): item.set(qn("w:"+key), value)
                borders.append(item)
            table._tbl.tblPr.append(borders)
            for index, source in enumerate(rows):
                row = table.add_row()
                props = row._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if index == 0: props.append(OxmlElement("w:tblHeader"))
                cells = [c for c in source.children if isinstance(c, Element)]
                for cell, node, width in zip(row.cells, cells, widths):
                    cell.width = Cm(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tc = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for edge, value in (("top",90),("bottom",90),("left",95),("right",95)):
                        val = OxmlElement("w:"+edge);val.set(qn("w:w"),str(value));val.set(qn("w:type"),"dxa");margins.append(val)
                    tc.append(margins)
                    if index == 0:
                        shade = OxmlElement("w:shd");shade.set(qn("w:fill"),"EAF0F6");tc.append(shade)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.15
                    if index == 0: p.paragraph_format.keep_with_next = True
                    inlines(p, node, bold=index == 0)
                    for run in p.runs: font(run, 9)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            for child in element.children: block(child)

    for element in body.children: block(element)
    # Word's default Title style may inherit a blue bottom rule. Remove it
    # explicitly rather than relying on a renderer's theme interpretation.
    for root in (doc.styles.element, doc.element):
        for border in list(root.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.save(destination)
    # Verify the output has editable paragraphs/tables and embedded images.
    check = Document(destination)
    assert len(check.tables) == 8 and len(check.inline_shapes) == 2
    assert any("6.3 使用过程中遇到的AI问题" in p.text for p in check.paragraphs)
    return len(check.paragraphs), len(check.tables)


def make_browser(source):
    style = re.search(r"<style>([\s\S]*?)</style>", source).group(1)
    body = re.search(r"<body>([\s\S]*?)</body>", source).group(1)
    extra = """body{max-width:none;margin:0;padding:0;background:#edf0f4}#report{background:white;max-width:880px;padding:32px 36px;margin:24px auto;outline:none;min-height:80vh}#report h1+p{text-align:center}#editor-bar{position:sticky;top:0;z-index:10;background:#fff;border-bottom:1px solid #c6ced8;padding:12px 24px;display:flex;gap:12px;align-items:center;flex-wrap:wrap;font-family:'Microsoft YaHei',sans-serif;font-size:14px}#editor-bar button{padding:7px 15px;border:1px solid #73889e;border-radius:4px;background:#f5f8fc;cursor:pointer}#edit-help{font-size:12px;color:#47566a}#status{color:#245287}@media print{body{background:white}#editor-bar{display:none!important}#report{padding:0;margin:0;max-width:none}}"""
    script = """(()=>{'use strict';const report=document.getElementById('report'),status=document.getElementById('status');document.getElementById('save-copy').onclick=()=>{const copy=document.documentElement.cloneNode(true);copy.querySelector('#status').textContent='副本已打开，可继续修改';const blob=new Blob(['<!doctype html>\\n'+copy.outerHTML],{type:'text/html;charset=utf-8'});const url=URL.createObjectURL(blob),link=document.createElement('a');link.href=url;link.download='2026_ECNU_PJ1_1_温正泽-我的修改.html';link.click();setTimeout(()=>URL.revokeObjectURL(url),1000);status.textContent='已生成保存副本，请查看下载位置';};document.getElementById('print-report').onclick=()=>print();report.addEventListener('input',()=>{status.textContent='有未保存修改，请点保存副本';});addEventListener('beforeunload',e=>{if(status.textContent.startsWith('有未保存')){e.preventDefault();e.returnValue='';}});})();"""
    bar = '<div id="editor-bar"><strong>报告编辑版</strong><button type="button" id="save-copy">保存副本</button><button type="button" id="print-report">打印或导出PDF</button><span id="status">点击正文即可修改</span><span id="edit-help">文字和表格可直接改；保存会下载新HTML，请用副本继续编辑。关闭前请保存。</span></div>'
    return '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>第一阶段报告 可编辑版</title><style>'+style+extra+'</style></head><body>'+bar+'<main id="report" contenteditable="true" spellcheck="false" aria-label="可编辑的报告正文">'+body+'</main><script>'+script+'</script></body></html>'


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--force',action='store_true');args=parser.parse_args()
    word=OUTPUT/(NAME+'-可编辑.docx');browser=OUTPUT/(NAME+'-网页编辑版.html')
    if not args.force and (word.exists() or browser.exists()):
        raise SystemExit('Editable copy exists; preserve user edits. Back up first before using --force.')
    source=(ROOT/'reports'/'第一阶段完成报告.html').read_text(encoding='utf-8')
    reader=Reader();reader.feed(source)
    html=next(c for c in reader.root.children if isinstance(c,Element) and c.tag=='html')
    body=next(c for c in html.children if isinstance(c,Element) and c.tag=='body')
    counts=make_word(body,word)
    browser.write_text(make_browser(source),encoding='utf-8')
    print(json.dumps({'word':str(word),'browser':str(browser),'paragraphs':counts[0],'editable_tables':counts[1]},ensure_ascii=False))


if __name__=='__main__':main()
